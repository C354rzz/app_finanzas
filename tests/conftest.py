from decimal import Decimal as D
from types import SimpleNamespace

import pytest
from django.contrib.auth import get_user_model

from apps.catalogos.models import Categoria, Concepto, Cuenta, Domicilio, Persona
from apps.catalogos.servicios import sembrar_catalogos
from apps.core.servicios import crear_hogar

# Renglones del presupuesto del Excel 2026 del usuario (montos y marcas; nombres ficticios).
# Marcas: F = fijo, T = con tarjeta, H = hormiga.
RENGLONES_EXCEL = {
    "Casa": [("200", "F"), ("100", ""), ("1500", "F"), ("1100", "F"), ("500", "F"),
             ("500", "F"), ("100", "F")],
    "Comida": [("3000", "F"), ("3000", "F"), ("1600", "F"), ("2000", "F"), ("600", "H"),
               ("400", "H")],
    "Familia": [("4641", "F"), ("800", "F")],
    "Transporte": [("1400", "F"), ("1000", "F")],
    "Deudas": [("1500", "F"), ("2000", "F")],
    "Salud": [("1000", "F")],
    "Suscripciones": [("239", "TH"), ("10", "T"), ("49", "T"), ("196", "F")],
    "Entretenimiento": [("500", "H")],
    "Otros": [("300", "F"), ("500", "TH"), ("500", "TH")],
}  # fmt: skip

# Encabezados de categoría del Excel: (texto, categoría, fila, columna del nombre).
BLOQUES_EXCEL = [
    ("🏡Casa", "Casa", 21, 3),
    ("🥑Comida", "Comida", 21, 9),
    ("❤️Familia", "Familia", 21, 15),
    ("🚓Transporte", "Transporte", 37, 3),
    ("✈️Viajes", "Viajes", 37, 9),
    ("🏦Deudas", "Deudas", 37, 15),
    ("🚑Salud", "Salud", 53, 3),
    ("📺Suscripciones", "Suscripciones", 53, 9),
    ("🏦Gastos anuales", "Gastos anuales", 53, 15),
    ("💅Cuidado personal", "Cuidado personal", 69, 3),
    ("📽️Entretenimiento", "Entretenimiento", 69, 9),
    ("🛸Otros", "Otros", 69, 15),
]


def _poner(hoja, fila, columna, *valores):
    for desplazamiento, valor in enumerate(valores):
        hoja.cell(fila, columna + desplazamiento, valor)


def _excel_ficticio(ruta, modificar=None):
    """Libro con el formato del «Financial Planner» y datos ficticios (de RENGLONES_EXCEL)."""
    from openpyxl import Workbook

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Presupuesto"
    hoja["C2"] = "Ingresos Promedio Mensual"
    _poner(hoja, 3, 3, "Salario mensual (neto)", "Fijo", None, None, 32977.52)
    _poner(hoja, 4, 3, "Bono", "Variable", None, None, 0)
    _poner(hoja, 5, 3, "-", "-")
    hoja["C14"], hoja["G14"] = "¿Qué % de tus ingresos quieres ahorrar?", 0.05
    for encabezado, categoria, fila, columna in BLOQUES_EXCEL:
        _poner(hoja, fila, columna, encabezado, "¿Gasto Fijo?", "¿Pago con tarjeta?",
               "Gasto hormiga 🐜", "Monto mensual")  # fmt: skip
        renglones = RENGLONES_EXCEL.get(categoria, [])
        for numero, (monto, marcas) in enumerate(renglones, start=1):
            _poner(hoja, fila + numero, columna, f"{categoria} {numero}", "F" in marcas,
                   "T" in marcas, "H" in marcas, float(monto))  # fmt: skip
        _poner(hoja, fila + len(renglones) + 1, columna, "-", False, False, False, 0)

    deudas = libro.create_sheet("Deudas")
    deudas["C1"] = "Tarjetas de Crédito"
    _poner(deudas, 2, 3, "Banco", "Tarjeta", "Tasa promedio*", "Saldo Actual", "Línea de crédito")
    formula = "=IFERROR(AVERAGEIFS('No borrar'!$H$2:$H$190,'No borrar'!$F$2:$F$190,C4),0)"
    _poner(deudas, 3, 3, "Banco Demo", "Oro", 0.45, 6839.01, 7100, None, "No")
    _poner(deudas, 4, 3, "Banco Demo", "Clásica", formula, 1000, 5000, None, "Sí")
    _poner(deudas, 5, 3, "Otro", "Tarjeta", formula, 0, 0)
    deudas["C15"] = "Créditos"
    _poner(deudas, 16, 3, "Tipo de crédito", "Deuda inicial", "Deuda actual", "Mensualidad")
    _poner(deudas, 17, 3, "Préstamo auto", 41130, 38679.72, 1500)
    _poner(deudas, 18, 3, "-", 0, 0, 0)

    metas = libro.create_sheet("Metas de Ahorro")
    metas["C11"] = "¿Cuál es tu ahorro actual?"
    for columna, nombre, actual, meses, tasa, meta in [
        (3, "🎁Regalo", 0, 6, 0.10, 3000),
        (7, "🏖️Vacaciones", 2000, 12, 0.05, 20000),
        (11, "🚗Auto", 0, 12, 0.10, 0),
        (15, "🎲Otros", None, None, None, None),
    ]:
        metas.cell(8, columna, nombre)
        for fila, valor in zip((11, 12, 13, 14), (actual, meses, tasa, meta), strict=True):
            metas.cell(fila, columna + 2, valor)

    patrimonio = libro.create_sheet("Patrimonio")
    patrimonio["C2"] = "Activos"
    _poner(patrimonio, 3, 3, "Activo", None, "Nombre", None, "Valor actual")
    activos = [
        ("-", None, None),
        ("🏡Casa/Departamento", "Casa ejemplo", 1500000),
        ("🚗Auto", "Auto ejemplo", 150000),
        ("💰Cuentas de Ahorro", "-", 25000),
        ("-", "-", 0),
    ]
    for fila, (tipo, nombre, valor) in enumerate(activos, start=4):
        _poner(patrimonio, fila, 3, tipo, None, nombre, None, valor)

    tasas = libro.create_sheet("No borrar")
    _poner(tasas, 1, 6, "Banco", "Tarjeta", "Tasa")
    _poner(tasas, 2, 6, "Banco Demo", "Oro", 0.6904399999999999)
    _poner(tasas, 3, 6, "Banco Demo", "Clásica", 0.73228)

    if modificar:
        modificar(libro)
    libro.save(ruta)
    return ruta


@pytest.fixture
def excel_ficticio(tmp_path):
    def crear(modificar=None):
        return _excel_ficticio(tmp_path / "planner.xlsx", modificar)

    return crear


@pytest.fixture
def usuario(db):
    return get_user_model().objects.create_user(
        email="julio@example.com", password="clave-segura-123"
    )


@pytest.fixture
def hogar(usuario):
    return crear_hogar("Familia Prueba", usuario)


@pytest.fixture
def otro_hogar(db):
    otro = get_user_model().objects.create_user(email="otro@example.com", password="clave-123-x")
    return crear_hogar("Otra Familia", otro)


@pytest.fixture
def cliente(client, usuario, hogar):
    """Cliente con sesión iniciada de un usuario que ya tiene hogar."""
    client.force_login(usuario)
    return client


@pytest.fixture
def catalogo(hogar):
    """Catálogo mínimo y ficticio de un hogar."""
    sembrar_catalogos(hogar)
    categorias = {c.nombre: c for c in Categoria.objects.del_hogar(hogar)}
    fidel = Domicilio.objects.create(hogar=hogar, alias="Casa Fidel")
    nomina = Cuenta.objects.create(hogar=hogar, nombre="Nómina", tipo=Cuenta.Tipo.DEBITO)
    return SimpleNamespace(
        categorias=categorias,
        monze=Persona.objects.create(hogar=hogar, nombre="Monze", parentesco="hija"),
        fidel=fidel,
        efectivo=Cuenta.objects.create(hogar=hogar, nombre="Efectivo", tipo=Cuenta.Tipo.EFECTIVO),
        nomina=nomina,
        tarjeta=Cuenta.objects.create(
            hogar=hogar,
            nombre="Tarjeta Oro",
            tipo=Cuenta.Tipo.CREDITO,
            linea_credito=D("7100"),
            saldo_actual=D("6839.01"),
            paga_total_mensual=False,
        ),
        prestamo=Cuenta.objects.create(
            hogar=hogar,
            nombre="Préstamo auto",
            tipo=Cuenta.Tipo.PRESTAMO,
            monto_inicial=D("41130"),
            mensualidad=D("1500"),
            saldo_actual=D("38679.72"),
        ),
        gasolina=Concepto.objects.create(
            hogar=hogar, categoria=categorias["Transporte"], nombre="Gasolina"
        ),
        recarga=Concepto.objects.create(
            hogar=hogar, categoria=categorias["Casa"], nombre="Recarga móvil"
        ),
        luz_fidel=Concepto.objects.create(
            hogar=hogar,
            categoria=categorias["Casa"],
            nombre="Luz",
            domicilio=fidel,
            cuenta=nomina,
            es_fijo=True,
        ),
    )


@pytest.fixture
def plantilla_excel(hogar):
    """Plantilla con los valores del Excel: ingresos 32,977.52 y gastos 29,235."""
    from apps.presupuesto.models import PlantillaGasto, PlantillaIngreso
    from apps.presupuesto.servicios import obtener_plantilla

    sembrar_catalogos(hogar)
    plantilla = obtener_plantilla(hogar)
    PlantillaIngreso.objects.create(
        hogar=hogar,
        plantilla=plantilla,
        nombre="Salario mensual (neto)",
        tipo_ingreso="salario",
        monto=D("32977.52"),
    )
    for nombre_categoria, renglones in RENGLONES_EXCEL.items():
        categoria = Categoria.objects.get(hogar=hogar, nombre=nombre_categoria)
        for numero, (monto, marcas) in enumerate(renglones, start=1):
            concepto = Concepto.objects.create(
                hogar=hogar, categoria=categoria, nombre=f"{nombre_categoria} {numero}"
            )
            PlantillaGasto.objects.create(
                hogar=hogar,
                plantilla=plantilla,
                concepto=concepto,
                monto=D(monto),
                es_fijo="F" in marcas,
                con_tarjeta="T" in marcas,
                es_hormiga="H" in marcas,
            )
    return plantilla


@pytest.fixture(autouse=True)
def media_temporal(settings, tmp_path):
    """Los archivos subidos en las pruebas van a una carpeta temporal."""
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture(autouse=True)
def respaldos_temporales(settings, tmp_path):
    """Los respaldos de las pruebas van a una carpeta temporal."""
    settings.RESPALDOS_DIR = tmp_path / "respaldos"


def _pdf_minimo(paginas):
    """PDF válido con una página por texto (Helvetica, WinAnsi). Solo para pruebas."""

    def escapar(linea):
        return linea.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objetos = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    hijos = []
    numero = 4
    for texto in paginas:
        lineas = " ".join(f"({escapar(linea)}) Tj T*" for linea in texto.split("\n"))
        flujo = f"BT /F1 10 Tf 12 TL 40 800 Td {lineas} ET".encode("latin-1")
        objetos[numero] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents " + f"{numero + 1} 0 R >>".encode()
        )
        objetos[numero + 1] = (
            f"<< /Length {len(flujo)} >>\nstream\n".encode() + flujo + b"\nendstream"
        )
        hijos.append(f"{numero} 0 R")
        numero += 2
    objetos[2] = f"<< /Type /Pages /Kids [{' '.join(hijos)}] /Count {len(hijos)} >>".encode()
    salida = bytearray(b"%PDF-1.4\n")
    posiciones = {}
    for clave in sorted(objetos):
        posiciones[clave] = len(salida)
        salida += f"{clave} 0 obj\n".encode() + objetos[clave] + b"\nendobj\n"
    inicio_xref = len(salida)
    total = max(objetos) + 1
    salida += f"xref\n0 {total}\n0000000000 65535 f \n".encode()
    for clave in range(1, total):
        salida += f"{posiciones[clave]:010d} 00000 n \n".encode()
    salida += (
        f"trailer\n<< /Size {total} /Root 1 0 R >>\nstartxref\n{inicio_xref}\n%%EOF\n".encode()
    )
    return bytes(salida)


@pytest.fixture
def crear_pdf():
    """Fábrica de PDFs ficticios: crear_pdf("texto página 1", "texto página 2")."""
    return lambda *paginas: _pdf_minimo(paginas or ("",))
