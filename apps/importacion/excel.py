"""Importación inicial desde el Excel «Financial Planner» (RF-ACC-02, RF-ACC-04).

Valida el formato completo antes de escribir y solo crea lo que falta: nunca modifica conceptos,
cuentas, metas, activos ni renglones que ya existen. La excepción es el catálogo global de tasas
de mercado, que se actualiza. Con aplicar=False, todo se revierte al final (simulación).
"""

import math
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.exceptions import InvalidFileException

from apps.catalogos.models import Categoria, Concepto, TasaMercado
from apps.catalogos.servicios import CATEGORIAS_INICIALES, sembrar_catalogos
from apps.importacion.clasificacion import normalizar
from apps.movimientos.models import TipoIngreso
from apps.presupuesto.models import Periodicidad, PlantillaGasto, PlantillaIngreso
from apps.presupuesto.servicios import obtener_plantilla

SECCIONES = [
    "tasas de mercado",
    "ingresos de la plantilla",
    "conceptos",
    "gastos de la plantilla",
    "tarjetas",
    "créditos",
    "metas",
    "activos",
]
HOJAS = ["Presupuesto", "Deudas", "Metas de Ahorro", "Patrimonio", "No borrar"]
# Celdas que identifican el formato de la plantilla: (hoja, celda, texto con el que empieza).
FORMATO = [
    ("Presupuesto", "C2", "Ingresos"),
    ("Presupuesto", "C14", "Que % de tus ingresos"),
    ("Deudas", "C1", "Tarjetas de credito"),
    ("Deudas", "C16", "Tipo de credito"),
    ("Metas de Ahorro", "C11", "Cual es tu ahorro actual"),
    ("Patrimonio", "C3", "Activo"),
    ("Patrimonio", "G3", "Valor actual"),
    ("No borrar", "F1", "Banco"),
    ("No borrar", "G1", "Tarjeta"),
    ("No borrar", "H1", "Tasa"),
]
# Presupuesto: encabezado de cada categoría (fila, columna del nombre), en el orden de
# CATEGORIAS_INICIALES. Cada bloque: 13 renglones con nombre, ¿fijo?, ¿tarjeta?, ¿hormiga?, monto.
BLOQUES = [(21, "C"), (21, "I"), (21, "O"), (37, "C"), (37, "I"), (37, "O"),
           (53, "C"), (53, "I"), (53, "O"), (69, "C"), (69, "I"), (69, "O")]  # fmt: skip
RENGLONES_POR_BLOQUE = 13
FILAS_INGRESOS = range(3, 11)
LIMITE = Decimal("1e10")
VACIOS = {"", "-"}
VERDADEROS = {"true", "verdadero", "si", "x"}


class ErrorExcel(Exception):
    """El archivo no es el Excel esperado; el mensaje es para el usuario."""


@dataclass
class Resumen:
    creados: Counter = field(default_factory=Counter)
    existentes: Counter = field(default_factory=Counter)
    avisos: list = field(default_factory=list)


def _columna(columna):
    return columna if isinstance(columna, int) else column_index_from_string(columna)


def _decimal(valor, decimales=2):
    """Número de la celda como Decimal, o None si no es un número utilizable."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, int | float):
        if not math.isfinite(valor):
            return None
        texto = str(valor)
    elif isinstance(valor, str):
        texto = valor.replace("$", "").replace(",", "").strip()
    else:
        return None
    try:
        numero = Decimal(texto).quantize(Decimal(1).scaleb(-decimales))
    except InvalidOperation:
        return None
    if not numero.is_finite():  # «NaN» escrito como texto
        return None
    return numero if abs(numero) < LIMITE else None


def _sin_icono(texto):
    """«🏡Casa» → «Casa»: quita emojis y signos del inicio."""
    return re.sub(r"^\W+", "", texto).strip()


class Hoja:
    """Valores calculados de una hoja y, aparte, si una celda tiene fórmula."""

    def __init__(self, valores, formulas):
        self.valores = valores
        self.formulas = formulas

    def valor(self, fila, columna):
        return self.valores.cell(fila, _columna(columna)).value

    def texto(self, fila, columna):
        valor = self.valor(fila, columna)
        texto = " ".join(str(valor).split()) if valor is not None else ""
        return "" if texto in VACIOS else texto

    def numero(self, fila, columna, decimales=2):
        return _decimal(self.valor(fila, columna), decimales)

    def bandera(self, fila, columna):
        valor = self.valor(fila, columna)
        return valor is True or (isinstance(valor, str) and normalizar(valor) in VERDADEROS)

    def es_formula(self, fila, columna):
        valor = self.formulas.cell(fila, _columna(columna)).value
        return isinstance(valor, str) and valor.startswith("=")


def _abrir(ruta):
    try:
        valores = load_workbook(ruta, data_only=True)
        formulas = load_workbook(ruta, data_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise ErrorExcel(
            "No se pudo abrir el archivo; debe ser el Excel (.xlsx) del planeador."
        ) from error
    faltantes = [nombre for nombre in HOJAS if nombre not in valores.sheetnames]
    if faltantes:
        raise ErrorExcel(f"Al Excel le falta la hoja «{faltantes[0]}».")
    return {nombre: Hoja(valores[nombre], formulas[nombre]) for nombre in HOJAS}


def _validar_formato(hojas):
    for nombre, celda, esperado in FORMATO:
        columna, fila = coordinate_from_string(celda)
        texto = normalizar(_sin_icono(hojas[nombre].texto(fila, columna)))
        if not texto.startswith(normalizar(esperado)):
            raise ErrorExcel(f"La hoja «{nombre}» no tiene el formato esperado (celda {celda}).")
    for (fila, columna), (categoria, _) in zip(BLOQUES, CATEGORIAS_INICIALES, strict=True):
        texto = normalizar(_sin_icono(hojas["Presupuesto"].texto(fila, columna)))
        if texto != normalizar(categoria):
            raise ErrorExcel(
                f"La hoja «Presupuesto» no tiene el formato esperado (celda {columna}{fila} "
                f"debería ser «{categoria}»)."
            )


@transaction.atomic
def importar_excel(ruta, hogar, aplicar=True):
    hojas = _abrir(ruta)
    _validar_formato(hojas)
    resumen = Resumen()
    sembrar_catalogos(hogar)
    _importar_tasas(hojas["No borrar"], resumen)
    _importar_presupuesto(hojas["Presupuesto"], hogar, resumen)
    if not aplicar:
        transaction.set_rollback(True)
    return resumen


def _importar_tasas(hoja, resumen):
    """Promedio por banco y tarjeta, como el AVERAGEIFS de la hoja Deudas del Excel."""
    tasas = defaultdict(list)
    for fila in range(2, hoja.valores.max_row + 1):
        institucion, producto = hoja.texto(fila, "F")[:80], hoja.texto(fila, "G")[:80]
        if not institucion and not producto:
            continue
        tasa = hoja.numero(fila, "H", decimales=6)
        if not institucion or not producto or tasa is None or not 0 <= tasa < 10:
            resumen.avisos.append(
                f"Tasas: el renglón {fila} no tiene banco, tarjeta o tasa válida."
            )
            continue
        tasas[(institucion, producto)].append(tasa)
    for (institucion, producto), valores in tasas.items():
        promedio = (sum(valores) / len(valores)).quantize(Decimal("0.0001"))
        _, nueva = TasaMercado.objects.update_or_create(
            institucion=institucion, producto=producto, defaults={"tasa_promedio": promedio}
        )
        (resumen.creados if nueva else resumen.existentes)["tasas de mercado"] += 1


def _monto(hoja, fila, columna, contexto, resumen):
    """Monto positivo de la celda; avisa si trae texto en lugar de un número."""
    valor = hoja.valor(fila, columna)
    monto = _decimal(valor)
    if monto is None and valor is not None and str(valor).strip() not in VACIOS:
        resumen.avisos.append(f"{contexto}: el monto «{valor}» no es un número; se omitió.")
    return monto if monto is not None and monto > 0 else None


def _importar_presupuesto(hoja, hogar, resumen):
    plantilla = obtener_plantilla(hogar)
    if not plantilla.ingresos.exists() and not plantilla.gastos.exists():
        porcentaje = hoja.numero(14, "G", decimales=4)
        if porcentaje is not None and 0 <= porcentaje <= 1:
            plantilla.porcentaje_ahorro = porcentaje
            plantilla.save(update_fields=["porcentaje_ahorro", "actualizado_en"])

    for fila in FILAS_INGRESOS:
        monto = _monto(hoja, fila, "G", f"Ingresos: renglón {fila}", resumen)
        if monto is None:
            continue
        nombre = (hoja.texto(fila, "C") or f"Ingreso {fila - 2}")[:80]
        if plantilla.ingresos.filter(nombre__iexact=nombre).exists():
            resumen.existentes["ingresos de la plantilla"] += 1
            continue
        PlantillaIngreso.objects.create(
            hogar=hogar,
            plantilla=plantilla,
            nombre=nombre,
            monto=monto,
            es_fijo=normalizar(hoja.texto(fila, "D")) != "variable",
            tipo_ingreso=(
                TipoIngreso.SALARIO if "salario" in normalizar(nombre) else TipoIngreso.OTRO
            ),
        )
        resumen.creados["ingresos de la plantilla"] += 1

    categorias = {c.nombre: c for c in Categoria.objects.del_hogar(hogar)}
    for (fila, columna), (nombre_categoria, _) in zip(BLOQUES, CATEGORIAS_INICIALES, strict=True):
        inicio = column_index_from_string(columna)
        vistos = set()
        for renglon in range(fila + 1, fila + 1 + RENGLONES_POR_BLOQUE):
            nombre = hoja.texto(renglon, inicio)[:80]
            if not nombre:
                continue
            if normalizar(nombre) in vistos:
                resumen.avisos.append(
                    f"{nombre_categoria}: «{nombre}» está repetido; se tomó el primero."
                )
                continue
            vistos.add(normalizar(nombre))
            es_fijo, con_tarjeta, es_hormiga = (
                hoja.bandera(renglon, inicio + desplazamiento) for desplazamiento in (1, 2, 3)
            )
            concepto = _concepto(
                hogar, categorias[nombre_categoria], nombre, es_fijo, es_hormiga, resumen
            )
            contexto = f"{nombre_categoria}: «{nombre}»"
            monto = _monto(hoja, renglon, inicio + 4, contexto, resumen)
            if monto is None:
                continue
            if PlantillaGasto.objects.filter(plantilla=plantilla, concepto=concepto).exists():
                resumen.existentes["gastos de la plantilla"] += 1
                continue
            PlantillaGasto.objects.create(
                hogar=hogar,
                plantilla=plantilla,
                concepto=concepto,
                monto=monto,
                es_fijo=es_fijo,
                con_tarjeta=con_tarjeta,
                es_hormiga=es_hormiga,
                periodicidad=(
                    Periodicidad.ANUAL
                    if nombre_categoria == "Gastos anuales"
                    else Periodicidad.MENSUAL
                ),
            )
            resumen.creados["gastos de la plantilla"] += 1


def _concepto(hogar, categoria, nombre, es_fijo, es_hormiga, resumen):
    existente = Concepto.objects.filter(
        hogar=hogar, categoria=categoria, nombre__iexact=nombre, domicilio=None
    ).first()
    if existente:
        resumen.existentes["conceptos"] += 1
        return existente
    resumen.creados["conceptos"] += 1
    return Concepto.objects.create(
        hogar=hogar, categoria=categoria, nombre=nombre, es_fijo=es_fijo, es_hormiga=es_hormiga
    )
